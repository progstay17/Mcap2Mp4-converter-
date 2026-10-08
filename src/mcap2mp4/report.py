"""Ekspor laporan CSV (BOM, pemisah koma/titik koma), JSON, XLSX (openpyxl), HTML (PRD FR-51..FR-54).

JSON memakai kode status/error netral bahasa (FR-82); CSV/XLSX/HTML memakai teks dari i18n.
"""
from __future__ import annotations

import csv
import html
import json
from datetime import datetime
from pathlib import Path

from .i18n import t

COLS = ["source", "status", "error", "message", "source_bytes", "elapsed_s", "mb_per_s", "outputs", "frames",
        "fps", "time_source", "duration_s", "resolution", "co64", "verify", "warnings", "sidecar"]


def _status(code, lang):
    return t(f"status.{code}", lang)


def _row(r: dict, lang) -> list:
    outs = r.get("outputs") or []
    err = r.get("error_code")
    verify = r.get("verify") or {}
    vtxt = ""
    if verify:
        vtxt = t("verify.ok", lang) if all(v.get("quick", True) and v.get("payload_matches_source", True) for v in verify.values()) else t("verify.failed", lang)
    return [r.get("source"), _status(r.get("status", ""), lang), t(f"err.{err}", lang) if err else "",
            r.get("error", ""), r.get("source_bytes"), r.get("elapsed_s"), r.get("mb_per_s"),
            "; ".join(Path(o["path"]).name for o in outs),
            "; ".join(f"{o.get('topic')}={o.get('frames')}" for o in outs if o.get("frames") is not None),
            r.get("fps_used"), r.get("time_source"),
            "; ".join(str(o.get("duration_s")) for o in outs if o.get("duration_s") is not None),
            "; ".join(f"{o['width']}x{o['height']}" for o in outs if o.get("width")),
            "; ".join(str(o.get("co64")) for o in outs if "co64" in o), vtxt,
            len(r.get("warnings") or []), r.get("sidecar", "")]


def _headers(lang):
    return [t(f"col.{c}", lang) for c in COLS]


def _summary_rows(summary, lang):
    c = summary["counts"]
    rows = [(t(f"status.{k}", lang), c.get(k, 0)) for k in ("success", "skipped", "incomplete", "out_of_sync", "failed")]
    rows += [(t("sum.app_version", lang), summary["app_version"]), (t("sum.files", lang), summary["files"]),
             (t("sum.total_bytes", lang), summary["total_source_bytes"]), (t("sum.elapsed", lang), summary["elapsed_s"]),
             (t("sum.throughput", lang), summary["throughput_mb_s"]), (t("sum.finished", lang), summary["finished_utc"])]
    return rows


def _warning_rows(results, lang):
    out = []
    for r in results:
        for w in r.get("warnings") or []:
            out.append([r.get("source"), w.get("topic", ""), w.get("code"), w.get("message")])
    return out


def _topic_rows(results):
    out = []
    for r in results:
        for topic, n in (r.get("topics") or {}).items():
            out.append([r.get("source"), topic, n])
    return out


def export(results, summary, out_dir, formats=("csv", "json"), delimiter=",", lang=None, stamp=None) -> list[Path]:
    lang = lang or summary.get("settings", {}).get("language") or "id"
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    paths = []
    for fmt in formats:
        p = out_dir / f"laporan_{stamp}.{fmt}"
        if fmt == "csv":
            with p.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f, delimiter=delimiter)
                w.writerow(_headers(lang))
                for r in results:
                    w.writerow(_row(r, lang))
        elif fmt == "json":
            slim = {k: v for k, v in summary.items()}
            p.write_text(json.dumps({"summary": slim, "files": results}, indent=2, ensure_ascii=False, default=str),
                         encoding="utf-8")
        elif fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = t("sheet.summary", lang)
            for row in _summary_rows(summary, lang):
                ws.append(list(row))
            ws = wb.create_sheet(t("sheet.files", lang))
            ws.append(_headers(lang))
            for r in results:
                ws.append(_row(r, lang))
            ws.freeze_panes = "A2"
            ws = wb.create_sheet(t("sheet.warnings", lang))
            ws.append([t("col.source", lang), t("col.topic", lang), "code", t("col.message", lang)])
            for row in _warning_rows(results, lang):
                ws.append(row)
            ws = wb.create_sheet(t("sheet.topics", lang))
            ws.append([t("col.source", lang), t("col.topic", lang), t("col.messages", lang)])
            for row in _topic_rows(results):
                ws.append(row)
            wb.save(p)
        elif fmt == "html":
            def tr(cells, tag="td"):
                return "<tr>" + "".join(f"<{tag}>{html.escape('' if c is None else str(c))}</{tag}>" for c in cells) + "</tr>"
            body = ["<h1>MCAP2MP4</h1><table border=1 cellpadding=4>"] + [tr(r) for r in _summary_rows(summary, lang)] + ["</table><br>"]
            body += ["<table border=1 cellpadding=4>", tr(_headers(lang), "th")] + [tr(_row(r, lang)) for r in results] + ["</table>"]
            p.write_text(f"<!doctype html><meta charset=utf-8><title>MCAP2MP4</title>{''.join(body)}", encoding="utf-8")
        else:
            continue
        paths.append(p)
    return paths
